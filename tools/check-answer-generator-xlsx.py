"""Inspect browser-exported XLSX files against their saved work documents."""
import json
import re
import sys
from pathlib import Path
from unicodedata import normalize
from zipfile import ZipFile

from openpyxl import load_workbook


def inspect(directory, name):
    work = json.loads((directory / f'{name}.json').read_text())
    with ZipFile(directory / f'{name}.xlsx') as archive:
        assert archive.testzip() is None, 'ZIP checksum failure'
    book = load_workbook(directory / f'{name}.xlsx')
    assert book.sheetnames == ['解答用紙', '模範解答']
    blank, model = book.worksheets
    for sheet in book:
        title = work['paper']['title']
        if work['paper']['writing'] == 'vertical':
            actual = ''.join(str(sheet.cell(r + 2, work['grid']['cols']).value or '') for r in range(len(title)))
            assert actual == title.replace('　', ''), (name, 'title changed during export')
        else:
            assert sheet.cell(1, 1).value == title, (name, 'title changed during export')
    for sheet in book:
        assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_A4
        assert sheet.page_setup.scale == 80
        assert sheet.page_setup.orientation == work['paper']['orientation']
        seen = set()
        for merged in sheet.merged_cells.ranges:
            assert merged.max_col <= work['grid']['cols']
            assert merged.max_row <= work['grid']['rows']
            for coordinate in merged.cells:
                assert coordinate not in seen, (name, sheet.title, merged, coordinate)
                seen.add(coordinate)
        expected_total = str(sum(field['points'] for field in work['fields']))
        values = [normalize('NFKC', str(cell.value)) for row in sheet for cell in row if cell.value is not None]
        assert expected_total in values, (name, sheet.title, expected_total)
    vertical = work['paper']['writing'] == 'vertical'
    if name == 'blanks':
        assert [(f['rects'][0]['w'], f['rects'][0]['h']) for f in work['fields']] == [(1, 20), (2, 8)]
    if name == 'kokugo':
        a, b = work['fields'][3:5]
        assert a['auto']['questionId'] == b['auto']['questionId']
        assert [work['texts'][f"{f['label']['c']},{f['label']['r']}"]['t'] for f in (a, b)] == ['（１）', '（２）']
        rect = work['fields'][8]['rects'][0]
        assert (rect['w'], rect['h']) == (1, 20)
    headings = []
    for key, text in work['texts'].items():
        if not text.get('centerAcross'):
            continue
        c, r = map(int, key.split(','))
        span = text['span']
        middle = (span - 1) // 2
        col, row = (c + middle + 1, r + 1) if vertical else (c + 1, r + middle * 2 + 1)
        headings.append((c, r, span))
        for sheet in book:
            cell = sheet.cell(row, col)
            assert cell.value == text['t'], (name, sheet.title, 'heading center', key)
            assert cell.alignment.horizontal == cell.alignment.vertical == 'center'
            if not vertical or span % 2 == 0:
                width = 2 if vertical else 1
                height = 1 if vertical else (4 if span % 2 == 0 else 2)
                assert any(m.min_col == col and m.min_row == row and m.max_col == col + width - 1
                           and m.max_row == row + height - 1 for m in sheet.merged_cells.ranges)
    for previous, current in zip(headings, headings[1:]):
        pc, pr, ps = previous
        cc, cr, cs = current
        assert cc + cs == pc - 1 if vertical else cr == pr + ps * 2 + 2
        for sheet in book:
            gap = (sheet.cell(r + 1, pc) for r in range(work['grid']['rows'])) if vertical else (
                sheet.cell(r + 1, c + 1) for r in range(pr + ps * 2, cr) for c in range(work['grid']['cols']))
            for cell in gap:
                assert cell.value is None
                assert all(getattr(cell.border, side).style is None for side in ('top', 'bottom', 'left', 'right'))
    for field in work['fields']:
        question_id = field['auto'].get('questionId')
        if question_id:
            parent = next(v for v in work['decorations'] if v['id'] == question_id)
            assert parent['questionHeading'] and parent['w'] == 1
            assert parent['h'] == (1 if vertical else 2)
            parent_text = work['texts'][f"{parent['c']},{parent['r']}"]['t']
            label = field['label']
            assert label['w'] == 1 and label['h'] == (1 if vertical else 2)
            suffix = work['texts'][f"{label['c']},{label['r']}"]['t']
            assert suffix.startswith(('(', '（')) and suffix.endswith((')', '）'))
            for sheet in book:
                cell = sheet.cell(parent['r'] + 1, parent['c'] + 1)
                assert normalize('NFKC', str(cell.value)) == normalize('NFKC', parent_text)
                assert cell.border.top.style == cell.border.left.style == 'thin'
                assert cell.alignment.horizontal == cell.alignment.vertical == 'center'
        label = field.get('label')
        if label:
            expected_label = work['texts'][f"{label['c']},{label['r']}"]['t']
            for sheet in book:
                actual_label = sheet.cell(label['r'] + 1, label['c'] + 1).value
                assert normalize('NFKC', str(actual_label)) == normalize('NFKC', expected_label)
        expected = work['answers'].get(field['id'], '')
        kind = field['auto']['kind']
        answer_cells = []
        for rect in field['rects']:
            assert rect['c'] >= 0 and rect['r'] >= 0
            assert rect['c'] + rect['w'] <= work['grid']['cols']
            assert rect['r'] + rect['h'] <= work['grid']['rows']
            row, col = rect['r'] + 1, rect['c'] + 1
            style = 'medium' if field['complete'] or field['type'] == 3 else field.get('line', 'thin')
            for sheet in book:
                assert sheet.cell(row, col).border.top.style == style
                assert sheet.cell(row, col).border.left.style == style
            if vertical:
                coordinates = [(c + 1, r + 1) for c in range(rect['c'] + rect['w'] - 1, rect['c'] - 1, -1)
                               for r in range(rect['r'], rect['r'] + rect['h'])]
            else:
                coordinates = [(c + 1, r + 1) for r in range(rect['r'], rect['r'] + rect['h'])
                               for c in range(rect['c'], rect['c'] + rect['w'])]
            for c, r in coordinates:
                cell = model.cell(r, c)
                if cell.value is None or kind == 'graph' and cell.value == 'O':
                    continue
                if kind == 'extract' and cell.value in ('初め', '終わり'):
                    continue
                answer_cells.append(cell)
                assert blank.cell(r, c).value is None, (name, field['id'], 'answer leaked into blank sheet')
        actual = ''.join(str(cell.value) for cell in answer_cells)
        if kind == 'extract':
            wave = re.split('[～〜~]', expected)
            expected = ''.join(part.strip() for part in wave) if len(wave) == 2 else expected.replace(' ', '').replace('　', '').replace('／', '').replace('/', '')
        assert actual == expected, (name, field['id'], actual, expected)
        for cell in answer_cells:
            if field['fixedChars']:
                assert len(str(cell.value)) == 1
                assert cell.font.sz == 11
            else:
                assert 6 <= cell.font.sz <= 11
                assert cell.alignment.horizontal == ('center' if vertical else 'left')
                assert cell.alignment.vertical == ('top' if vertical else 'center')
                if vertical and len(str(cell.value)) > 1 and any(ord(ch) > 127 for ch in str(cell.value)):
                    assert cell.alignment.textRotation == 255
        if kind == 'chars' and not field['exact'] and field['chars'] >= 5:
            cells = []
            for rect in field['rects']:
                if vertical:
                    cells.extend((c, r) for c in range(rect['c'] + rect['w'] - 1, rect['c'] - 1, -1)
                                 for r in range(rect['r'], rect['r'] + rect['h'], 2))
                else:
                    cells.extend((c, r) for r in range(rect['r'], rect['r'] + rect['h'], 2)
                                 for c in range(rect['c'], rect['c'] + rect['w']))
            c, r = cells[field['chars'] * 8 // 10]
            border = blank.cell(r + 1, c + 1).border
            assert (border.top if vertical else border.left).style == 'mediumDashed'
    print(f'PASS: {name} Excel: answers, scoring, fonts, borders, 80% line, A4, heading centers, section gaps and separate parent/subquestion cells')


if __name__ == '__main__':
    directory = Path(sys.argv[1])
    for name in ('vertical', 'horizontal', 'kokugo', 'sections', 'blanks'):
        inspect(directory, name)
